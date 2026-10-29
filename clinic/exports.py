# exports.py - helpers for the admin screens: date ranges and Excel downloads.
from datetime import datetime, timedelta, timezone as dt_timezone

from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError

# A download never holds more than this many rows (keeps the server safe)
MAX_EXPORT_ROWS = 50000


# Reads ?date_from=2026-10-01&date_to=2026-10-28 (both optional, both included in the range).
# Returns (date_from, date_to) as date objects or None.
def read_date_range(request):
    values = []
    for name in ("date_from", "date_to"):
        raw = request.query_params.get(name, "").strip()
        if not raw:
            values.append(None)
            continue
        day = parse_date(raw)
        if day is None:
            raise ValidationError({name: "Use the date format YYYY-MM-DD."})
        values.append(day)
    date_from, date_to = values
    if date_from and date_to and date_from > date_to:
        raise ValidationError("The start date must not be after the end date.")
    return date_from, date_to


# Keeps only the rows whose date column (e.g. "created_at") falls inside the range.
# For a date-and-time column, __date uses the project's time zone (Africa/Nairobi),
# so "today" means today in Kenya. A plain date column (like visit_date) needs no __date.
def filter_by_range(rows, field, date_from, date_to, has_time=True):
    suffix = "__date" if has_time else ""
    if date_from:
        rows = rows.filter(**{field + suffix + "__gte": date_from})
    if date_to:
        rows = rows.filter(**{field + suffix + "__lte": date_to})
    return rows


# A cell that starts with = + - @ could be run as a formula when opened in Excel.
# Some of our text comes from outside (e.g. the email typed in a failed sign-in),
# so we put a ' in front to make Excel show it as plain text.
def safe_cell(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


# A date-time as Kenya local time, without a time zone (Excel cannot store time zones)
def local_naive(moment):
    if moment is None:
        return None
    return timezone.localtime(moment).replace(tzinfo=None)


# Builds an .xlsx file and returns it as a download.
#   sheet_name - the tab name inside the workbook
#   headers    - the column titles
#   rows       - an iterable of lists, one list per row
#   filename   - what the browser saves it as
#   info       - short lines like "Period: 1 Oct 2026 to 28 Oct 2026" put on a second tab
def excel_response(sheet_name, headers, rows, filename, info=()):
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="0F766E")        # the teal of the site
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    widths = [len(str(h)) for h in headers]
    count = 0
    for row in rows:
        if count >= MAX_EXPORT_ROWS:
            break
        clean = [safe_cell(v) for v in row]
        ws.append(clean)
        for i, v in enumerate(clean):
            text = v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime) else str(v if v is not None else "")
            widths[i] = max(widths[i], min(len(text), 60))
        count += 1
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width + 2
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
    if count:
        ws.auto_filter.ref = ws.dimensions

    # a second tab that says what this file contains
    about = wb.create_sheet("About this file")
    about.append(["Exported", local_naive(timezone.now())])
    about.append(["Rows", count])
    about.column_dimensions["A"].width = 14
    about.column_dimensions["B"].width = 60
    for key, value in info:
        about.append([key, safe_cell(value)])
    if count >= MAX_EXPORT_ROWS:
        about.append(["Note", f"Only the first {MAX_EXPORT_ROWS} rows are included. Narrow the dates to see the rest."])
    about["B1"].number_format = "yyyy-mm-dd hh:mm"

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
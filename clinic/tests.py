import io
import shutil
import tempfile
from pathlib import Path
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from PIL import Image
from rest_framework.test import APITestCase

from accounts.models import Role, User
from clinic.models import Lab, SystemLog

# Uploaded test images go into a temporary folder, not into your real media/ folder
TEMP_MEDIA = tempfile.mkdtemp()

# A made-up CNN answer, so the tests do not need TensorFlow or the real model
FAKE_PREDICTION = {
    "label": "parasitized",
    "probability": 0.93,
    "threshold": 0.19,
    "model_version": "test",
    "warning": "",
}


# Makes a small PNG image in memory
def make_png(size=(130, 130)):
    buf = io.BytesIO()
    Image.new("RGB", size, (150, 60, 120)).save(buf, format="PNG")
    buf.seek(0)
    buf.name = "cell.png"
    return buf


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class SystemTests(APITestCase):

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)

    # Runs before EACH test: a fresh lab and one user per role
    def setUp(self):
        self.lab = Lab.objects.create(name="Test Lab", location="Kisumu")
        self.nurse = self.make_user("nurse@test.org", Role.NURSE)
        self.tech = self.make_user("tech@test.org", Role.LAB_TECHNICIAN)
        self.doctor = self.make_user("doctor@test.org", Role.DOCTOR)
        self.admin = self.make_user("admin@test.org", Role.ADMIN)

    def make_user(self, email, role):
        return User.objects.create_user(
            email=email, password="StrongPass#123", name=role, role=role, lab=self.lab
        )

    # Pretend to be this user from now on
    def act_as(self, user):
        self.client.force_authenticate(user=user)

    # The nurse registers a patient and a visit. Returns the visit id
    def new_visit(self):
        self.act_as(self.nurse)
        patient = self.client.post("/api/patients/", {"name": "Jane Achieng", "age": 30}, format="json")
        self.assertEqual(patient.status_code, 201)
        visit = self.client.post("/api/visits/", {
            "patient": patient.json()["id"], "lab": self.lab.pk,
            "temperature": "38.5", "heart_rate": 90, "blood_pressure": "120/80", "symptoms": "fever",
        }, format="json")
        self.assertEqual(visit.status_code, 201)
        self.assertEqual(visit.json()["status"], "sent_to_lab")
        return visit.json()["id"]

    # The technician opens a report for the visit. Returns the report id
    def new_report(self, visit_id):
        self.act_as(self.tech)
        report = self.client.post("/api/lab-reports/", {"visit": visit_id, "notes": "smear"}, format="json")
        self.assertEqual(report.status_code, 201)
        return report.json()["id"]

    # ---------- 1. The whole nurse -> lab -> doctor flow ----------
    @patch("clinic.services.ml.predict", return_value=FAKE_PREDICTION)
    def test_full_workflow(self, _fake):
        visit_id = self.new_visit()
        report_id = self.new_report(visit_id)

        # Technician uploads an image and the (fake) CNN answers
        upload = self.client.post(f"/api/lab-reports/{report_id}/images/", {"image": make_png()}, format="multipart")
        self.assertEqual(upload.status_code, 201)
        self.assertEqual(upload.json()["predicted_label"], "parasitized")

        # Technician sends the report to the doctor
        self.assertEqual(self.client.post(f"/api/lab-reports/{report_id}/refer/").status_code, 200)

        # Doctor records the diagnosis
        self.act_as(self.doctor)
        diagnosis = self.client.post(
            f"/api/visits/{visit_id}/diagnosis/",
            {"result": "positive", "agrees_with_prediction": True}, format="json",
        )
        self.assertEqual(diagnosis.status_code, 201)
        self.assertEqual(self.client.get(f"/api/visits/{visit_id}/").json()["status"], "diagnosed")

        # A second diagnosis for the same visit is refused
        again = self.client.post(
            f"/api/visits/{visit_id}/diagnosis/",
            {"result": "negative", "agrees_with_prediction": False}, format="json",
        )
        self.assertEqual(again.status_code, 400)

        # The audit log recorded the diagnosis
        self.assertTrue(SystemLog.objects.filter(action="CREATE_DIAGNOSIS").exists())

    # ---------- 2. A report with no images cannot be sent ----------
    def test_cannot_refer_without_images(self):
        report_id = self.new_report(self.new_visit())
        self.assertEqual(self.client.post(f"/api/lab-reports/{report_id}/refer/").status_code, 400)

    # ---------- 3. The doctor cannot diagnose before the lab has sent the case ----------
    def test_doctor_cannot_diagnose_too_early(self):
        visit_id = self.new_visit()
        self.act_as(self.doctor)
        response = self.client.post(
            f"/api/visits/{visit_id}/diagnosis/",
            {"result": "positive", "agrees_with_prediction": True}, format="json",
        )
        self.assertEqual(response.status_code, 400)

    # ---------- 4. Roles are enforced ----------
    def test_roles_are_enforced(self):
        self.act_as(self.tech)
        self.assertEqual(self.client.post("/api/visits/", {}, format="json").status_code, 403)
        self.act_as(self.doctor)
        self.assertEqual(self.client.post("/api/patients/", {"name": "X", "age": 5}, format="json").status_code, 403)
        self.act_as(self.nurse)
        self.assertEqual(self.client.get("/api/logs/").status_code, 403)
        self.assertEqual(self.client.get("/api/users/").status_code, 403)

    # ---------- 5. Not logged in = no access ----------
    def test_login_required(self):
        self.client.force_authenticate(user=None)
        self.assertEqual(self.client.get("/api/patients/").status_code, 401)

    # ---------- 6. A nurse sees only her own visits ----------
    def test_nurse_sees_only_own_visits(self):
        self.new_visit()
        self.act_as(self.nurse)
        self.assertEqual(self.client.get("/api/visits/").json()["count"], 1)
        other_nurse = self.make_user("nurse2@test.org", Role.NURSE)
        self.act_as(other_nurse)
        self.assertEqual(self.client.get("/api/visits/").json()["count"], 0)

    # ---------- 7. Bad uploads are rejected ----------
    @patch("clinic.services.ml.predict", return_value=FAKE_PREDICTION)
    def test_bad_uploads_rejected(self, _fake):
        report_id = self.new_report(self.new_visit())
        url = f"/api/lab-reports/{report_id}/images/"
        not_an_image = SimpleUploadedFile("a.png", b"not an image", content_type="image/png")
        self.assertEqual(self.client.post(url, {"image": not_an_image}, format="multipart").status_code, 400)
        too_small = make_png(size=(10, 10))
        self.assertEqual(self.client.post(url, {"image": too_small}, format="multipart").status_code, 400)

    # ---------- 8. If the model file is missing the answer is 503, not a crash ----------
    def test_missing_model_gives_503(self):
        report_id = self.new_report(self.new_visit())
        with self.settings(MODEL_PATH=Path(TEMP_MEDIA) / "missing.h5"):
            response = self.client.post(
                f"/api/lab-reports/{report_id}/images/", {"image": make_png()}, format="multipart"
            )
        self.assertEqual(response.status_code, 503)

    # ---------- 9. Login rules ----------
    def test_login_and_deactivation(self):
        self.client.force_authenticate(user=None)
        good = {"email": "nurse@test.org", "password": "StrongPass#123"}
        bad = {"email": "nurse@test.org", "password": "wrong"}
        self.assertEqual(self.client.post("/api/auth/login/", good, format="json").status_code, 200)
        self.assertEqual(self.client.post("/api/auth/login/", bad, format="json").status_code, 400)
        self.assertTrue(SystemLog.objects.filter(action="LOGIN_FAILED").exists())   # failed attempt recorded
        # A deactivated account cannot log in even with the right password
        self.nurse.is_active = False
        self.nurse.save()
        self.assertEqual(self.client.post("/api/auth/login/", good, format="json").status_code, 400)
    # ---------- Correcting a diagnosis ----------
    # Makes a visit that the doctor has already diagnosed. Returns the visit id
    def diagnosed_visit(self, result="positive", agrees=True):
        visit_id = self.new_visit()
        report_id = self.new_report(visit_id)
        self.client.post(f"/api/lab-reports/{report_id}/images/", {"image": make_png()}, format="multipart")
        self.client.post(f"/api/lab-reports/{report_id}/refer/")
        self.act_as(self.doctor)
        self.client.post(
            f"/api/visits/{visit_id}/diagnosis/",
            {"result": result, "agrees_with_prediction": agrees, "notes": "first note"}, format="json",
        )
        return visit_id

    @patch("clinic.services.ml.predict", return_value=FAKE_PREDICTION)
    def test_doctor_can_correct_a_diagnosis_and_history_is_kept(self, _fake):
        visit_id = self.diagnosed_visit(result="positive", agrees=True)

        fix = self.client.post(
            f"/api/visits/{visit_id}/correct-diagnosis/",
            {"result": "negative", "agrees_with_prediction": False, "notes": "re-read the smear", "reason": "Slide was mislabelled"},
            format="json",
        )
        self.assertEqual(fix.status_code, 200)
        body = fix.json()
        self.assertEqual(body["result"], "negative")
        self.assertTrue(body["overruled"])
        self.assertEqual(len(body["corrections"]), 1)
        self.assertEqual(body["corrections"][0]["previous_result"], "positive")
        self.assertEqual(body["corrections"][0]["reason"], "Slide was mislabelled")

        # the visit now shows the corrected diagnosis, with the history
        visit = self.client.get(f"/api/visits/{visit_id}/").json()
        self.assertEqual(visit["diagnosis_result"], "negative")
        self.assertEqual(len(visit["diagnosis"]["corrections"]), 1)
        # and the audit log recorded it
        self.assertTrue(SystemLog.objects.filter(action="CORRECT_DIAGNOSIS").exists())

    @patch("clinic.services.ml.predict", return_value=FAKE_PREDICTION)
    def test_correction_needs_a_reason_and_a_real_change(self, _fake):
        visit_id = self.diagnosed_visit(result="positive", agrees=True)
        url = f"/api/visits/{visit_id}/correct-diagnosis/"
        base = {"result": "negative", "agrees_with_prediction": False, "notes": "x"}
        # no reason
        self.assertEqual(self.client.post(url, {**base, "reason": "  "}, format="json").status_code, 400)
        # nothing changed
        same = {"result": "positive", "agrees_with_prediction": True, "notes": "first note", "reason": "typo"}
        self.assertEqual(self.client.post(url, same, format="json").status_code, 400)

    @patch("clinic.services.ml.predict", return_value=FAKE_PREDICTION)
    def test_only_a_doctor_can_correct_and_only_a_diagnosed_visit(self, _fake):
        visit_id = self.diagnosed_visit()
        url = f"/api/visits/{visit_id}/correct-diagnosis/"
        body = {"result": "negative", "agrees_with_prediction": False, "notes": "", "reason": "test"}
        for user in (self.nurse, self.tech, self.admin):
            self.act_as(user)
            self.assertEqual(self.client.post(url, body, format="json").status_code, 403)
        # a visit with no diagnosis yet
        fresh = self.new_visit()
        self.act_as(self.doctor)
        self.assertEqual(
            self.client.post(f"/api/visits/{fresh}/correct-diagnosis/", body, format="json").status_code, 400
        )

    # ---------- admin: labs and staff ----------
    def test_admin_can_list_and_add_labs_others_cannot(self):
        for user in (self.nurse, self.tech, self.doctor):
            self.act_as(user)
            self.assertEqual(self.client.get("/api/labs/").status_code, 403)
        self.act_as(self.admin)
        made = self.client.post("/api/labs/", {"name": "Nyanza Lab", "location": "Kisumu"}, format="json")
        self.assertEqual(made.status_code, 201)
        # the same name again (any capitals) is refused
        again = self.client.post("/api/labs/", {"name": "nyanza lab", "location": "x"}, format="json")
        self.assertEqual(again.status_code, 400)
        names = [lab["name"] for lab in self.client.get("/api/labs/").json()]
        self.assertEqual(names, ["Nyanza Lab", "Test Lab"])
        self.assertTrue(SystemLog.objects.filter(action="CREATE_LAB").exists())

    def test_staff_list_can_be_filtered_and_shows_lab_name(self):
        self.act_as(self.admin)
        rows = self.client.get("/api/users/?role=doctor").json()["results"]
        self.assertEqual([r["email"] for r in rows], ["doctor@test.org"])
        self.assertEqual(rows[0]["lab_name"], "Test Lab")
        found = self.client.get("/api/users/?search=TECH@").json()["results"]
        self.assertEqual([r["email"] for r in found], ["tech@test.org"])

    def test_admin_cannot_lock_themselves_out(self):
        self.act_as(self.admin)
        url = f"/api/users/{self.admin.pk}/"
        self.assertEqual(self.client.patch(url, {"is_active": False}, format="json").status_code, 400)
        self.assertEqual(self.client.patch(url, {"role": "nurse"}, format="json").status_code, 400)
        # but they can still change someone else, and rename themselves
        other = f"/api/users/{self.nurse.pk}/"
        self.assertEqual(self.client.patch(other, {"is_active": False}, format="json").status_code, 200)
        self.assertEqual(self.client.patch(url, {"name": "Chief Admin"}, format="json").status_code, 200)

    # ---------- admin: timeframe and Excel downloads ----------
    def test_timeframe_filters_staff_and_logs(self):
        from datetime import timedelta
        from django.utils import timezone
        old = timezone.now() - timedelta(days=40)
        User.objects.filter(pk=self.nurse.pk).update(created_at=old)
        SystemLog.objects.create(actor=self.nurse, actor_role="nurse", action="LOGIN")
        SystemLog.objects.filter(action="LOGIN").update(logged_at=old)
        SystemLog.objects.create(actor=self.doctor, actor_role="doctor", action="LOGOUT")
        self.act_as(self.admin)
        since = (timezone.localdate() - timedelta(days=7)).isoformat()
        emails = [u["email"] for u in self.client.get(f"/api/users/?date_from={since}").json()["results"]]
        self.assertNotIn("nurse@test.org", emails)
        self.assertIn("doctor@test.org", emails)
        actions = [r["action"] for r in self.client.get(f"/api/logs/?date_from={since}").json()["results"]]
        self.assertIn("LOGOUT", actions)
        self.assertNotIn("LOGIN", actions)
        # bad or reversed dates are refused
        self.assertEqual(self.client.get("/api/logs/?date_from=yesterday").status_code, 400)
        self.assertEqual(self.client.get("/api/logs/?date_from=2026-10-10&date_to=2026-10-01").status_code, 400)

    def test_excel_downloads_follow_the_filters_and_are_admin_only(self):
        from openpyxl import load_workbook
        SystemLog.objects.create(actor=None, action="LOGIN_FAILED", description="=cmd|calc")
        for user in (self.nurse, self.tech, self.doctor):
            self.act_as(user)
            self.assertEqual(self.client.get("/api/users/export/").status_code, 403)
            self.assertEqual(self.client.get("/api/logs/export/").status_code, 403)
        self.act_as(self.admin)
        staff = self.client.get("/api/users/export/?role=doctor")
        self.assertEqual(staff.status_code, 200)
        self.assertIn("spreadsheetml", staff["Content-Type"])
        rows = list(load_workbook(io.BytesIO(staff.content))["Staff"].values)
        self.assertEqual(rows[0][0], "Name")
        self.assertEqual([r[1] for r in rows[1:]], ["doctor@test.org"])
        logs = self.client.get("/api/logs/export/?action=LOGIN_FAILED")
        sheet = list(load_workbook(io.BytesIO(logs.content))["Activity logs"].values)
        self.assertEqual(len(sheet), 2)
        self.assertEqual(sheet[1][1], "Not signed in")
        self.assertTrue(sheet[1][4].startswith("'"))        # text that looks like a formula is made safe
        self.assertTrue(SystemLog.objects.filter(action="EXPORT_LOGS").exists())

    # ---------- patients: edit ----------
    def test_nurse_and_admin_can_edit_a_patient_others_cannot(self):
        self.new_visit()                                  # makes "Jane Achieng"
        pid = self.client.get("/api/patients/").json()["results"][0]["id"]
        url = f"/api/patients/{pid}/"
        self.act_as(self.nurse)
        fixed = self.client.patch(url, {"name": "Jane A. Achieng", "age": 31}, format="json")
        self.assertEqual(fixed.status_code, 200)
        self.assertEqual(fixed.json()["name"], "Jane A. Achieng")
        self.assertEqual(self.client.patch(url, {"age": 500}, format="json").status_code, 400)
        self.assertTrue(SystemLog.objects.filter(action="UPDATE_PATIENT").exists())
        for user in (self.tech, self.doctor):
            self.act_as(user)
            self.assertEqual(self.client.patch(url, {"age": 20}, format="json").status_code, 403)
            self.assertEqual(self.client.get(url).status_code, 200)       # but they may look

    # ---------- visits: filters and the admin records export ----------
    @patch("clinic.services.ml.predict", return_value=FAKE_PREDICTION)
    def test_visit_filters_and_records_export(self, _fake):
        from openpyxl import load_workbook
        visit_id = self.diagnosed_visit(result="negative", agrees=False)     # doctor overruled the scan
        other = self.new_visit()                                             # a second, undiagnosed visit
        patient_id = self.client.get(f"/api/visits/{visit_id}/").json()["patient"]["id"]
        self.act_as(self.admin)
        mine = self.client.get(f"/api/visits/?patient={patient_id}").json()["count"]
        self.assertEqual(mine, 1)                      # each new_visit() registers a new patient
        self.assertEqual(self.client.get("/api/visits/?overruled=1").json()["count"], 1)
        self.assertEqual(self.client.get("/api/visits/?result=negative").json()["count"], 1)
        self.assertEqual(self.client.get("/api/visits/?result=positive").json()["count"], 0)
        self.assertEqual(self.client.get("/api/visits/?search=achieng").json()["count"], 2)
        self.assertEqual(self.client.get("/api/visits/?search=nobody").json()["count"], 0)
        self.assertEqual(self.client.get("/api/visits/?date_from=2999-01-01").json()["count"], 0)
        # only the admin may download
        for user in (self.nurse, self.tech, self.doctor):
            self.act_as(user)
            self.assertEqual(self.client.get("/api/visits/export/").status_code, 403)
        self.act_as(self.admin)
        file = self.client.get("/api/visits/export/?overruled=1")
        self.assertEqual(file.status_code, 200)
        rows = list(load_workbook(io.BytesIO(file.content))["Records"].values)
        self.assertEqual(len(rows), 2)                                       # header + the one overruled visit
        header = rows[0]
        row = dict(zip(header, rows[1]))
        self.assertEqual(row["Patient"], "Jane Achieng")
        self.assertEqual(row["Scan says"], "positive")
        self.assertEqual(row["Diagnosis"], "negative")
        self.assertEqual(row["Overruled scan"], "Yes")
        self.assertTrue(SystemLog.objects.filter(action="EXPORT_RECORDS").exists())
        self.assertIsNotNone(other)

    # ---------- passwords ----------
    def test_password_reset_by_email_link_and_change_password(self):
        from django.core import mail
        self.client.force_authenticate(user=None)
        # unknown email gets the same answer and no mail
        self.assertEqual(self.client.post("/api/auth/password-reset/", {"email": "ghost@test.org"}, format="json").status_code, 200)
        self.assertEqual(len(mail.outbox), 0)
        self.assertEqual(self.client.post("/api/auth/password-reset/", {"email": "nurse@test.org"}, format="json").status_code, 200)
        self.assertEqual(len(mail.outbox), 1)
        link = [w for w in mail.outbox[0].body.split() if "reset-password?" in w][0]
        query = dict(part.split("=") for part in link.split("?")[1].split("&"))
        done = self.client.post("/api/auth/password-reset/confirm/",
                                {"uid": query["uid"], "token": query["token"], "new_password": "Brand#New2026pw"}, format="json")
        self.assertEqual(done.status_code, 200)
        self.assertTrue(User.objects.get(email="nurse@test.org").check_password("Brand#New2026pw"))
        # the same link cannot be used twice
        again = self.client.post("/api/auth/password-reset/confirm/",
                                 {"uid": query["uid"], "token": query["token"], "new_password": "Another#Pass2026"}, format="json")
        self.assertEqual(again.status_code, 400)
        # a signed-in user changes their own password
        self.act_as(self.doctor)
        wrong = self.client.post("/api/auth/change-password/", {"old_password": "nope", "new_password": "Fresh#Pass2026"}, format="json")
        self.assertEqual(wrong.status_code, 400)
        ok = self.client.post("/api/auth/change-password/", {"old_password": "StrongPass#123", "new_password": "Fresh#Pass2026"}, format="json")
        self.assertEqual(ok.status_code, 200)